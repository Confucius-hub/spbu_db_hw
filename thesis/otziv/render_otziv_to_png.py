#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Рендер отзыва (.docx) в PNG-превью формата A4 через matplotlib."""
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from docx import Document
import textwrap, os

DOCX = '/home/user/spbu_db_hw/thesis/otziv/Отзыв_о_практике_Байханов.docx'
OUT  = '/home/user/spbu_db_hw/thesis/otziv/Отзыв_превью.png'
plt.rcParams['font.family'] = 'Liberation Serif'

# A4 portrait, inches
PW, PH = 8.27, 11.69
ML, MR, MT = 1.18, 0.59, 0.79
TW = PW - ML - MR          # text width

doc = Document(DOCX)

fig = plt.figure(figsize=(PW, PH), dpi=150)
ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, PW); ax.set_ylim(0, PH)
ax.invert_yaxis(); ax.axis('off')
ax.add_patch(Rectangle((0, 0), PW, PH, facecolor='white', edgecolor='none'))

y = MT


def emit(text, size=14, bold=False, italic=False, align='justify',
         indent=0.0, gap_before=0.0, gap_after=0.10):
    global y
    y += gap_before
    char_w = size * 0.48 / 72.0
    maxc = max(8, int((TW - indent) / char_w))
    lines = textwrap.wrap(text, width=maxc) or ['']
    line_h = size / 72.0 * 1.4
    for i, ln in enumerate(lines):
        x0 = ML + (indent if i == 0 else 0)
        if align == 'center':
            ax.text(PW / 2, y, ln, fontsize=size, ha='center', va='top',
                    fontweight='bold' if bold else 'normal',
                    fontstyle='italic' if italic else 'normal')
        else:
            ax.text(x0, y, ln, fontsize=size, ha='left', va='top',
                    fontweight='bold' if bold else 'normal',
                    fontstyle='italic' if italic else 'normal')
        y += line_h
    y += gap_after


# Заголовок
emit('Отзыв о прохождении', size=17, align='center', gap_after=0.02)
emit('учебной практики', size=17, align='center', gap_after=0.18)

# Таблица
tbl = doc.tables[0]
rows = [(r.cells[0].text, r.cells[1].text) for r in tbl.rows]
tbl_x, tbl_w, col1 = ML, TW, 1.4
row_h = 0.32
for k, (a, b) in enumerate(rows):
    ry = y + k * row_h
    ax.add_patch(Rectangle((tbl_x, ry), tbl_w, row_h, fill=False,
                           edgecolor='black', linewidth=0.8))
    ax.add_patch(Rectangle((tbl_x, ry), col1, row_h, fill=False,
                           edgecolor='black', linewidth=0.8))
    ax.text(tbl_x + 0.08, ry + row_h / 2, a, fontsize=11, va='center',
            fontstyle='italic')
    ax.text(tbl_x + col1 + 0.1, ry + row_h / 2, b, fontsize=11, va='center',
            fontweight='bold' if k == 0 else 'normal',
            fontstyle='normal' if k == 0 else 'italic')
y += len(rows) * row_h + 0.22

# Тело
paras = [p for p in doc.paragraphs if p.text.strip()]
# пропускаем 2 строки заголовка
body = paras[2:]
for p in body:
    txt = p.text.strip()
    is_bullet = p.style.name == 'List Bullet'
    is_sign = txt.startswith('Руководитель практики')
    if is_bullet:
        emit('•  ' + txt, size=13, indent=0.5, gap_after=0.04)
    elif is_sign:
        y += 0.3
        emit(txt, size=12, align='left', gap_after=0.0)
    else:
        emit(txt, size=14, indent=0.5, gap_after=0.12)

fig.savefig(OUT, dpi=150)
plt.close()
print('Saved →', OUT)
