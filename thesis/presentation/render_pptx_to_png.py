#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Render the actual .pptx slides to PNG via matplotlib (reads real shapes)."""
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
import io, os, textwrap
from PIL import Image
import numpy as np

PPTX   = '/home/user/spbu_db_hw/thesis/presentation/Презентация_ВКР_2026.pptx'
OUTDIR = '/home/user/spbu_db_hw/thesis/presentation/slides_preview_v2'
os.makedirs(OUTDIR, exist_ok=True)
plt.rcParams['font.family'] = 'Liberation Serif'

prs = Presentation(PPTX)
SW = prs.slide_width  / 914400.0
SH = prs.slide_height / 914400.0


def emu(v):
    return (v or 0) / 914400.0


def fill_hex(shape):
    try:
        if shape.fill.type == 1:
            c = shape.fill.fore_color.rgb
            return '#%02x%02x%02x' % (c[0], c[1], c[2])
    except Exception:
        pass
    return None


def color_hex(run):
    try:
        if run.font.color and run.font.color.type is not None:
            c = run.font.color.rgb
            return '#%02x%02x%02x' % (c[0], c[1], c[2])
    except Exception:
        pass
    return '#000000'


for idx, slide in enumerate(prs.slides, 1):
    fig = plt.figure(figsize=(SW, SH), dpi=150)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, SW)
    ax.set_ylim(0, SH)
    ax.invert_yaxis()
    ax.axis('off')
    ax.add_patch(Rectangle((0, 0), SW, SH, facecolor='white', edgecolor='none', zorder=0))

    for shape in slide.shapes:
        x, y = emu(shape.left), emu(shape.top)
        w, h = emu(shape.width), emu(shape.height)

        # picture
        if shape.shape_type == 13:
            try:
                img = Image.open(io.BytesIO(shape.image.blob)).convert('RGB')
                ax.imshow(np.array(img), extent=[x, x + w, y + h, y],
                          aspect='auto', zorder=2)
            except Exception as e:
                print('  img err', e)
            continue

        # rectangle fill
        fh = fill_hex(shape)
        if fh:
            ax.add_patch(Rectangle((x, y), w, h, facecolor=fh,
                                   edgecolor='none', zorder=1))

        # text
        if shape.has_text_frame:
            cur_y = y + 0.06
            for para in shape.text_frame.paragraphs:
                runs = para.runs
                if not runs:
                    cur_y += 0.16
                    continue
                txt = ''.join(r.text for r in runs)
                if not txt.strip():
                    cur_y += 0.13
                    continue
                r0 = runs[0]
                sz = r0.font.size.pt if r0.font.size else 10
                bold = bool(r0.font.bold)
                col = color_hex(r0)
                align = para.alignment
                if align == PP_ALIGN.CENTER:
                    tx, ha = x + w / 2, 'center'
                elif align == PP_ALIGN.RIGHT:
                    tx, ha = x + w - 0.06, 'right'
                else:
                    tx, ha = x + 0.06, 'left'
                char_w = sz * 0.46 / 72.0
                maxc = max(4, int((w - 0.12) / char_w))
                lines = []
                for seg in txt.split('\n'):
                    lines.extend(textwrap.wrap(seg, width=maxc) or [''])
                line_h = sz / 72.0 * 1.32
                for ln in lines:
                    cur_y += line_h * 0.5
                    ax.text(tx, cur_y, ln, fontsize=sz, color=col,
                            fontweight='bold' if bold else 'normal',
                            ha=ha, va='center', zorder=3)
                    cur_y += line_h * 0.55

    fig.savefig(f'{OUTDIR}/slide_{idx:02d}.png', dpi=150)
    plt.close()
    print(f'slide_{idx:02d}.png')

print('ALL PREVIEWS DONE →', OUTDIR)
