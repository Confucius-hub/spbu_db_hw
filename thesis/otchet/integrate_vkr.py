# -*- coding: utf-8 -*-
"""Integrate the missing port figures into the user's existing thesis (.docx),
add new section 3.4, renumber 3.4/3.5 -> 3.5/3.6, fix cross-refs, and replace
the manual table of contents with a clickable auto-updating Word TOC field."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = '/tmp/user_report_fixed.docx'
FIG = '/home/user/spbu_db_hw/thesis/otchet/figures'
OUTDIR = '/home/user/spbu_db_hw/thesis/vkr'
os.makedirs(OUTDIR, exist_ok=True)
OUT = f'{OUTDIR}/ВКР_интегрированная.docx'

doc = Document(SRC)
P = doc.paragraphs

def get_style(name):
    for s in doc.styles:
        if s.name == name:
            return s
    raise KeyError(name)

# ── heading paragraph indices (from inspection) ──
CH = [76, 105, 171, 213, 265, 277]                       # Heading 1
SUB = [107,116,130,141,156, 173,184,195,207,
       219,230,242,248,257]                              # Heading 2
# capture key objects BEFORE mutating
anchor   = P[248]      # "3.4. Ограничения..."  -> insertion point + renumber to 3.5
p_vyvody = P[257]      # "3.5. Выводы по главе 3" -> 3.6
p_cr1    = P[218]      # cross-ref "...3.4 «Ограничения»"
p_cr2    = P[238]      # cross-ref "...3.4 настоящей работы"
p_zakl   = P[265]      # ЗАКЛЮЧЕНИЕ (insert вывод-bullet before it)
toc_head = P[53]       # СОДЕРЖАНИЕ
manual_toc = [P[i] for i in range(55, 75)]   # manual TOC entry lines

# ════════════════ A. Configure & apply heading styles ════════════════
def cfg_heading(style_name, size, align):
    st = get_style(style_name)
    st.font.name = 'Times New Roman'; st.font.size = Pt(size)
    st.font.bold = True; st.font.color.rgb = RGBColor(0,0,0)
    rpr = st.element.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), 'Times New Roman')
    pf = st.paragraph_format
    pf.alignment = align; pf.space_before = Pt(12); pf.space_after = Pt(6)
    pf.keep_with_next = True; pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(0)
cfg_heading('Heading 1', 14, WD_ALIGN_PARAGRAPH.CENTER)
cfg_heading('Heading 2', 14, WD_ALIGN_PARAGRAPH.LEFT)

def apply_style(par, style, align):
    par.style = get_style(style)
    par.alignment = align
    par.paragraph_format.first_line_indent = Cm(0)

for i in CH:  apply_style(P[i], 'Heading 1', WD_ALIGN_PARAGRAPH.CENTER)
for i in SUB: apply_style(P[i], 'Heading 2', WD_ALIGN_PARAGRAPH.LEFT)

# ════════════════ B. Renumber + cross-refs (before insertion) ════════════════
def replace_in_runs(par, old, new):
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new); return True
    # fallback: rebuild from full text
    full = par.text
    if old in full:
        for r in list(par.runs): r.text = ''
        par.runs[0].text = full.replace(old, new); return True
    return False

replace_in_runs(anchor,   '3.4.', '3.5.')   # heading Ограничения
replace_in_runs(p_vyvody, '3.5.', '3.6.')   # heading Выводы по главе 3
replace_in_runs(p_cr1,    '3.4', '3.5')     # cross-ref (only one 3.4 in this para)
replace_in_runs(p_cr2,    '3.4', '3.5')

# ════════════════ C. Build new section 3.4 (insert before anchor) ════════════════
def ins_heading(text):
    p = anchor.insert_paragraph_before(); p.style = get_style('Heading 2')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(text); r.bold = True; r.font.name='Times New Roman'; r.font.size=Pt(14)
    return p

def ins_body(text):
    p = anchor.insert_paragraph_before()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(1.25); pf.space_after = Pt(0)
    r = p.add_run(text); r.font.name='Times New Roman'; r.font.size=Pt(14)
    return p

def ins_figure(path, width_cm=16.0):
    p = anchor.insert_paragraph_before(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(path, width=Cm(width_cm))
    return p

def ins_caption(text):
    p = anchor.insert_paragraph_before(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text); r.font.name='Times New Roman'; r.font.size=Pt(12)
    return p

def ins_table_caption(text):
    p = anchor.insert_paragraph_before(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.space_before = Pt(8)
    r = p.add_run(text); r.font.name='Times New Roman'; r.font.size=Pt(12)
    return p

def _add_borders(t):
    tblPr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4')
        e.set(qn('w:space'),'0'); e.set(qn('w:color'),'000000')
        borders.append(e)
    tblPr.append(borders)

def ins_table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))   # appended at end
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _add_borders(t)
    for i,h in enumerate(headers):
        c = t.rows[0].cells[i]; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.paragraphs[0].add_run(h); r.bold=True; r.font.name='Times New Roman'; r.font.size=Pt(12)
    for row in rows:
        cells = t.add_row().cells
        for i,v in enumerate(row):
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cells[i].paragraphs[0].add_run(str(v)); r.font.name='Times New Roman'; r.font.size=Pt(12)
    anchor._p.addprevious(t._tbl)   # move to insertion point
    return t

# --- content ---
ins_heading('3.4. Качественный анализ радиолокационных сигнатур ложных целей '
            'на акваториях арктических портов')

ins_body('Количественные результаты переноса методики (раздел 3.2) дополнены '
  'качественным визуальным анализом радиолокационных сцен трёх арктических '
  'портов, представляющих полный спектр типов ложных целей. Принципиально, что '
  'за период 2022–2025 гг. на этих акваториях не зафиксировано подтверждённых '
  'нефтяных разливов, поэтому любая выделенная тёмная зона заведомо является '
  'ложной целью — это формирует естественный контрольный полигон (нулевую '
  'гипотезу) для оценки специфичности детектирования. На каждой сцене применён '
  'единый детектор: текстурная маска гладкой поверхности (локальное '
  'среднеквадратичное отклонение) отделяет воду и лёд от шероховатой суши, после '
  'чего внутри маски применяется адаптивный порог T = μ − k·σ с морфологической '
  'фильтрацией.')

ins_body('Порт Печенга (Кольский полуостров, Баренцево море) представляет '
  'безлёдный тип ложных целей (рисунок 3.3). Зимой 2026 г. в защищённых от ветра '
  'участках залива формируются ветровые тени — зоны пониженного волнения с низким '
  'коэффициентом обратного рассеяния; летом 2025 г. к ним добавляются биогенные '
  'плёнки и штилевые области. Обе сигнатуры по яркости неотличимы от нефтяного слика.')
ins_figure(f'{FIG}/fig_pechenga.png')
ins_caption('Рисунок 3.3 — Порт Печенга (Sentinel-1, VV): а) февраль 2026 г. (зима); '
  'б) август 2025 г. (лето); в, г) детекция тёмных зон. Ложные цели — ветровые '
  'тени и биогенные/штилевые плёнки')

ins_body('Нефтеналивной терминал Варандей (Печорское море, 68,8° с. ш.) '
  'иллюстрирует ложные цели ледового происхождения (рисунок 3.4). Суша '
  '(Малоземельская тундра) расположена в левой части кадра и даёт яркий отклик, '
  'открытое море — в правой. В конце мая 2025 г. (весенний ледоход) тёмные зоны '
  'порождаются открытой водой и ветровыми тенями; в середине апреля 2025 г. '
  '(поздняя зима) — гладким ровным первогодним льдом, радиолокационный отклик '
  'которого имитирует нефтяную плёнку.')
ins_figure(f'{FIG}/fig_varandey.png')
ins_caption('Рисунок 3.4 — Терминал Варандей (Sentinel-1, VV): а) конец мая 2025 г. '
  '(ледоход); б) середина апреля 2025 г. (зима); в, г) детекция тёмных зон. Слева — '
  'тундра (суша), справа — Печорское море; ледовые и ветровые ложные цели')

ins_body('Порт Сабетта (завод «Ямал СПГ», Обская губа Карского моря, 71,3° с. ш.) — '
  'наиболее сложный ледовый полигон (рисунок 3.5). Полуостров Ямал и порт '
  'расположены слева, акватория губы — справа. В октябре 2022 г. (начало '
  'ледостава) формируется жировой и ниласовый лёд — сильнейший двойник нефти; в '
  'декабре 2025 г. тёмные зоны образованы гладким припаем и свежезамёрзшими '
  'разводьями. Яркие объекты в нижней части декабрьской сцены соответствуют '
  'инфраструктуре порта и торосистому льду вдоль судового канала, а не нефти.')
ins_figure(f'{FIG}/fig_sabetta.png')
ins_caption('Рисунок 3.5 — Порт Сабетта (Sentinel-1, VV): а) октябрь 2022 г. '
  '(ледостав); б) конец декабря 2025 г. (зимний лёд); в, г) детекция тёмных зон. '
  'Слева — п-ов Ямал и порт, справа — Обская губа; ледовые ложные цели')

ins_table_caption('Таблица 3.3 — Преобладающие типы ложных целей на исследованных акваториях')
ins_table(['Порт', 'Акватория', 'Преобладающий тип ложных целей'],
  [['Печенга', 'Баренцево море', 'Ветровые тени, биогенные плёнки (безлёдный тип)'],
   ['Варандей', 'Печорское море', 'Сезонный первогодний лёд, ветровые тени'],
   ['Сабетта', 'Карское море (Обская губа)', 'Жировой/ниласовый лёд, припай']])
ins_body('')  # spacer

ins_body('Дополнительным подтверждением ограничений радиолокационного '
  'детектирования служит крупнейшая арктическая нефтяная катастрофа — разлив '
  '≈21 тыс. тонн дизельного топлива при аварии ТЭЦ-3 в Норильске 29 мая 2020 г. '
  '(рисунок 3.6). Топливо прошло по цепочке Далдыкан → Амбарная → оз. Пясино. '
  'Однако оперативный мониторинг разлива вёлся преимущественно средствами '
  'оптической съёмки Sentinel-2: тонкая летучая плёнка дизельного топлива на '
  'узких реках в период весеннего ледохода не формирует устойчивой сигнатуры в '
  'C-диапазоне, а наблюдаемый на радарных сценах сигнал определяется сезонной '
  'динамикой «лёд → открытая вода» (доля тёмных пикселей изменилась с 7,6 % до '
  '4,9 % за 12 суток). Этот пример обосновывает фокус разработанной системы на '
  'морских акваториях портов, где применима физика SAR-детектирования, и '
  'необходимость явного выделения льда в отдельный класс.')
ins_figure(f'{FIG}/fig_norilsk.png')
ins_caption('Рисунок 3.6 — Бассейн оз. Пясино (район разлива НТЭК, Норильск) на '
  'снимках Sentinel-1 IW GRD (RGB-композит R = σ⁰_VV, G = σ⁰_VH, B = VV/VH): '
  'а) 3 июня 2020 г.; б) 15 июня 2020 г.')

ins_body('Качественные наблюдения настоящего раздела непосредственно объясняют '
  'отмеченный в разделе 3.2 рост доли ложноположительных срабатываний при '
  'переносе модели на акватории с активным ледообразованием и подтверждают '
  'ключевое проектное решение работы — выделение начальных форм льда в отдельный '
  'класс семантической сегментации.')

# ════════════════ D. Add a bullet to "Выводы по главе 3" ════════════════
new_vyvod = p_zakl.insert_paragraph_before()
pf = new_vyvod.paragraph_format
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
pf.first_line_indent = Cm(1.25)
r = new_vyvod.add_run('Качественный анализ радиолокационных сцен трёх арктических '
  'портов (Печенга, Варандей, Сабетта), не содержавших подтверждённых разливов '
  'в 2022–2025 гг., визуально подтвердил основной тезис работы: тёмное пятно в '
  'SAR является необходимым, но не достаточным признаком нефти, а ветровые тени, '
  'биогенные плёнки и начальные формы льда формируют неотличимые ложные цели '
  '(раздел 3.4).')
r.font.name='Times New Roman'; r.font.size=Pt(14)

# ════════════════ E. Replace manual TOC with auto field ════════════════
# insert TOC field paragraph right after СОДЕРЖАНИЕ heading
toc_p = toc_head.insert_paragraph_before()  # placeholder, will move after head
# actually build the field paragraph and place after toc_head
toc_field_p = OxmlElement('w:p')
toc_head._p.addnext(toc_field_p)
# build runs: begin field, instrText, separate, placeholder text, end
def _run_with(child):
    r = OxmlElement('w:r'); r.append(child); return r
beginChar = OxmlElement('w:fldChar'); beginChar.set(qn('w:fldCharType'),'begin')
instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve')
instr.text = 'TOC \\o "1-3" \\h \\z \\u'
sepChar = OxmlElement('w:fldChar'); sepChar.set(qn('w:fldCharType'),'separate')
ph_t = OxmlElement('w:t'); ph_t.text = 'Содержание обновится автоматически при открытии в Word (или: правый клик → «Обновить поле» / F9).'
endChar = OxmlElement('w:fldChar'); endChar.set(qn('w:fldCharType'),'end')
for child in (beginChar, instr, sepChar):
    toc_field_p.append(_run_with(child))
toc_field_p.append(_run_with(ph_t))
toc_field_p.append(_run_with(endChar))
# remove the stray placeholder paragraph we created
toc_p._p.getparent().remove(toc_p._p)
# remove manual TOC entry paragraphs
for p in manual_toc:
    p._p.getparent().remove(p._p)

# ════════════════ F. Tell Word to update fields on open ════════════════
settings = doc.settings.element
if settings.find(qn('w:updateFields')) is None:
    upd = OxmlElement('w:updateFields'); upd.set(qn('w:val'),'true')
    settings.insert(0, upd)

doc.save(OUT)
print('SAVED', OUT)
